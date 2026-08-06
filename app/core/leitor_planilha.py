from pathlib import Path
from openpyxl import load_workbook
import sys
##import os


def carregar_mapa_pacotes():
    
    ##caminhos = []

    if getattr(sys, "frozen", False):
        pasta_base = Path(sys._MEIPASS)
    # Executando pelo .exe
        ##pasta_base = Path(sys.executable).parent
        ##caminhos.append(Path(sys.executable).parent / "DePara.xlsx")
        
    # Durante o desenvolvimento
    ##caminhos.append(Path(__file__).parent.parent.parent / "DePara.xlsx")
    
    else: 
        pasta_base = Path(__file__).parent.parent.parent
        
    caminho = pasta_base / "DePara.xlsx"
    #for caminho in caminhos:
        
         #print(f"Procurando planilha em: {caminho}")
        
    
         #if caminho.exists():
             
    wb = load_workbook(caminho, data_only=True)
    ws = wb.active
            
    mapa = {}

    ##else:
    # Executando pelo VS Code
        ##pasta_base = Path(__file__).parent.parent.parent

    ##caminho = pasta_base / "DePara.xlsx"

    

    for origem, destino in ws.iter_rows(
                min_row=2,
                max_col=2,
                values_only=True
            ):

                if origem is None or destino is None:
                    continue

                mapa[str(origem).strip()] = str(destino).strip()

    return mapa

#raise FileNotFoundError(
    #    "Não foi possível localizar o arquivo DePara.xlsx"
   # )
    

# ===== TESTE =====

##if __name__ == "__main__":

    #mapa = carregar_mapa_pacotes()

    #print(f"Total de registros: {len(mapa)}")

    #for codigo in [
     #   "351",
      #  "25701351",
       # "25702351",
        #"25703351"
    #]:
     #   print(f"{codigo} -> {mapa.get(codigo)}")